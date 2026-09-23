"""Builds the complete thesis as ONE Word file (Hairfall_Thesis.docx) in the format of the syllabus
(A4; margins L 3 cm, R/T/B 2.5 cm; Times New Roman 12 pt; line spacing 1.5; captions per syllabus).

All formatting is done with Word STYLES, so the font, size or spacing can be changed later in one place
(Home > Styles). Run:  python build_thesis.py
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from PIL import Image

sys.path.insert(0, os.path.dirname(__file__))
from omml import latex_to_omml  # noqa: E402

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(BASE), "Hairfall_Thesis.docx")
FONT = "Times New Roman"

TITLE = ("Comparative Benchmarking of CatBoost, TabPFN, and TabFM for Explainable "
         "Multi-Tier Hair Fall Risk Stratification")
AUTHOR = "Nirajan Shahi"
REGNO = "Roll No. 49/079"
SUPERVISOR = "Asst. Prof. Jagadish Bhatta"
DEPT = "Central Department of Computer Science and Information Technology"
UNIV = "Tribhuvan University"
DEGREE = "Master of Science in Computer Science and Information Technology (M.Sc. CSIT)"
DATE = "[Month, Year]"

refs = json.load(open(os.path.join(BASE, "refs.json"), encoding="utf-8"))

# --------------------------------------------------------------------------------------------
# 1. Read the sources and number citations (by first appearance) and equations
# --------------------------------------------------------------------------------------------
chapters = [open(os.path.join(BASE, f"ch{i}.md"), encoding="utf-8").read() for i in range(1, 6)]
appendix_src = open(os.path.join(BASE, "appendices.md"), encoding="utf-8").read()

cite_no, order = {}, []
for txt in chapters + [appendix_src]:
    for m in re.finditer(r"\[@(\w+)\]", txt):
        k = m.group(1)
        if k not in cite_no:
            order.append(k)
            cite_no[k] = len(order)
missing = [k for k in order if k not in refs]
assert not missing, missing

eq_no = {}
for txt in chapters:
    for m in re.finditer(r"^EQ\[(\w+)\]:", txt, re.M):
        eq_no[m.group(1)] = len(eq_no) + 1


def resolve(text):
    text = re.sub(r"\[@(\w+)\]", lambda m: f"[{cite_no[m.group(1)]}]", text)
    text = re.sub(r"\{eq:(\w+)\}", lambda m: str(eq_no[m.group(1)]), text)
    return text


# --------------------------------------------------------------------------------------------
# 2. Document, page setup and styles
# --------------------------------------------------------------------------------------------
doc = Document()


def set_style_font(style, size=12, bold=False, italic=False, color=RGBColor(0, 0, 0), name=FONT):
    f = style.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in list(rfonts.attrib):
        del rfonts.attrib[a]
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{a}"), name)


def para_fmt(style, align=None, before=0, after=6, line=1.5, keep_next=False, page_break=False, left=None,
             first=None):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.page_break_before = page_break
    pf.widow_control = True
    if align is not None:
        pf.alignment = align
    if left is not None:
        pf.left_indent = left
    if first is not None:
        pf.first_line_indent = first


styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, 12)
para_fmt(normal, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6)

h1 = styles["Heading 1"]
set_style_font(h1, 14, bold=True)
para_fmt(h1, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12, keep_next=True, page_break=True)
h2 = styles["Heading 2"]
set_style_font(h2, 12, bold=True)
para_fmt(h2, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, keep_next=True)
h3 = styles["Heading 3"]
set_style_font(h3, 12, bold=True)
para_fmt(h3, align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=6, keep_next=True)
h4 = styles["Heading 4"]
set_style_font(h4, 12, bold=False, italic=True)
para_fmt(h4, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=6, keep_next=True)


def new_style(name, base="Normal", size=12, bold=False, italic=False, align=None, before=0, after=6, line=1.5,
              keep_next=False, page_break=False, left=None, first=None, font=FONT):
    st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = styles[base]
    set_style_font(st, size, bold, italic, name=font)
    para_fmt(st, align=align, before=before, after=after, line=line, keep_next=keep_next, page_break=page_break,
             left=left, first=first)
    return st


front_head = new_style("Front Heading", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12,
                       keep_next=True, page_break=True)
app_head = new_style("Appendix Heading", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12,
                     keep_next=True, page_break=True)
for st in (front_head, app_head):
    ppr = st.element.get_or_add_pPr()
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "0")
    ppr.append(lvl)
new_style("Subhead", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, before=6)
new_style("Table Caption", align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, before=6, after=4, line=1.0)
new_style("Figure Caption", align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=10, line=1.0)
new_style("Table Text", size=10, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
new_style("Equation", align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=4, line=1.5)
new_style("List Item", left=Cm(1.0), first=Cm(-0.6), after=3)
new_style("Code", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0, font="Courier New")
new_style("Reference", align=WD_ALIGN_PARAGRAPH.LEFT, left=Cm(1.0), first=Cm(-1.0), after=6, line=1.15)
new_style("Cover 16", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
new_style("Cover 16 Bold", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
new_style("Cover 14", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
new_style("Cover Title", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
new_style("Blank 14", size=14, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
new_style("Front Body", align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=Cm(1.27))
new_style("Front Plain", align=WD_ALIGN_PARAGRAPH.LEFT)
new_style("Abbrev", align=WD_ALIGN_PARAGRAPH.LEFT, left=Cm(2.5), first=Cm(-2.5), after=2)
new_style("Contents Entry", align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.5)
for i in (1, 2, 3):
    try:
        st = styles[f"TOC {i}"]
    except KeyError:
        st = styles.add_style(f"TOC {i}", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
    set_style_font(st, 12, bold=(i == 1))
    para_fmt(st, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.5, left=Cm(0.6 * (i - 1)))

sec = doc.sections[0]


def page_setup(s):
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(3.0)
    s.right_margin = Cm(2.5)
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.header_distance = Cm(1.25)
    s.footer_distance = Cm(1.25)


page_setup(sec)

# --------------------------------------------------------------------------------------------
# 3. Helpers: runs, fields, placeholders, equations, tables, figures
# --------------------------------------------------------------------------------------------
PLACEHOLDER = re.compile(r"\[[^\]]*(?:CHECK|Month|Year|Date|DD|name|Name|version|Insert|confirm|VERIFY|Logo|logo|"
                         r"Co-supervisor|examiner|stamp)[^\]]*\]")
SUBSCRIPT = re.compile(r"\b(h|F|φ|TP|FP|FN|d|D|x|p)_(\{[^}]+\}|[A-Za-z0-9]+)(?![A-Za-z_])")


def add_text(p, text, bold=None, italic=None):
    """Adds text with placeholder highlighting and math-style subscripts."""
    text = resolve(text)
    pos = 0
    pieces = []
    for m in PLACEHOLDER.finditer(text):
        pieces.append((text[pos:m.start()], False))
        pieces.append((m.group(0), True))
        pos = m.end()
    pieces.append((text[pos:], False))
    for chunk, hl in pieces:
        if not chunk:
            continue
        spos = 0
        for m in SUBSCRIPT.finditer(chunk):
            if m.start() > spos:
                _run(p, chunk[spos:m.start()], hl, bold, italic)
            _run(p, m.group(1), hl, bold, True)
            r = _run(p, m.group(2).strip("{}"), hl, bold, italic)
            r.font.subscript = True
            spos = m.end()
        if spos < len(chunk):
            _run(p, chunk[spos:], hl, bold, italic)


def _run(p, text, hl=False, bold=None, italic=None):
    r = p.add_run(text)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if hl:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def para(text="", style="Normal", bold=None, italic=None, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text, bold, italic)
    if align is not None:
        p.alignment = align
    return p


def blank(n=1, style="Blank 14"):
    for _ in range(n):
        doc.add_paragraph("", style=style)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def fld_char(run, kind):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._r.append(el)


def instr(run, code):
    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = code
    run._r.append(el)


def add_field_paragraph_simple(p, code, placeholder_text="1"):
    r = p.add_run()
    fld_char(r, "begin")
    r2 = p.add_run()
    instr(r2, code)
    r3 = p.add_run()
    fld_char(r3, "separate")
    p.add_run(placeholder_text)
    r5 = p.add_run()
    fld_char(r5, "end")


def add_footer_page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field_paragraph_simple(p, " PAGE ")
    for r in p.runs:
        r.font.name = FONT
        r.font.size = Pt(12)


def clear_footer(section):
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        for r in list(p.runs):
            r._r.getparent().remove(r._r)


def set_pgnum(section, fmt, start=1):
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    pg_mar = sect_pr.find(qn("w:pgMar"))
    pg_mar.addnext(pg)  # schema order: pgSz, pgMar, ..., pgNumType, cols


def add_equation(latex, number):
    p = doc.add_paragraph(style="Equation")
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Cm(7.75), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    p._p.append(parse_xml(latex_to_omml(latex)))
    p.add_run(f"\t({number})")
    return p


def add_figure(path, caption):
    full = os.path.normpath(os.path.join(BASE, path))
    with Image.open(full) as im:
        w, h = im.size
    width = 15.0
    height = width * h / w
    if height > 13.5:
        height = 13.5
        width = height * w / h
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(full, width=Cm(width))
    cp = doc.add_paragraph(style="Figure Caption")
    add_text(cp, caption)
    return cp


def set_cell_borders_and_width(cell, width_cm):
    cell.width = Cm(width_cm)


def add_table(rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # column widths proportional to content length
    lens = []
    for c in range(ncols):
        col = [header[c]] + [r[c] for r in body if c < len(r)]
        lens.append(max(6, min(45, max(len(resolve(x)) for x in col))))
    weights = [x ** 0.85 for x in lens]
    total = 15.5
    widths = [total * w / sum(weights) for w in weights]
    size = 8.5 if ncols >= 7 else 10
    for ri, row in enumerate([header] + body):
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            cell.width = Cm(widths[ci])
            p = cell.paragraphs[0]
            p.style = styles["Table Text"]
            txt = row[ci] if ci < len(row) else ""
            add_text(p, txt, bold=(ri == 0) or None)
            for r in p.runs:
                r.font.size = Pt(size)
    # repeat header row
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for row in t.rows:  # do not split rows across pages
        trpr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trpr.append(cs)
    doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(4)
    return t


def toc_field(entries, code, first_style_for=None):
    """Adds a TOC-type field, pre-filled with entries [(level, text)], so it also reads well before updating."""
    paras = []
    for lvl, text in entries:
        p = doc.add_paragraph(style=f"TOC {min(lvl,3)}" if lvl else "Contents Entry")
        p.add_run(resolve(text))
        paras.append(p)
    if not paras:
        return
    first, last = paras[0], paras[-1]
    r0 = first.runs[0]._r
    begin = OxmlElement("w:r")
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); begin.append(b)
    ins = OxmlElement("w:r")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code; ins.append(it)
    sep = OxmlElement("w:r")
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); sep.append(s)
    r0.addprevious(begin); r0.addprevious(ins); r0.addprevious(sep)
    end = OxmlElement("w:r")
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); end.append(e)
    last._p.append(end)


# --------------------------------------------------------------------------------------------
# 4. Parse chapter sources into blocks
# --------------------------------------------------------------------------------------------
HEAD = re.compile(r"^(\d+(?:\.\d+)+)\s+([A-Z].{1,110})$")
APPHEAD = re.compile(r"^([A-D]\.\d+)\s+([A-Z].{1,110})$")
TABCAP = re.compile(r"^Table ([A-Z0-9]+\.\d+):")
IMG = re.compile(r"^!\[(.*)\]\((.+)\)\s*$")
SUBHEADS = {"General objective", "Specific objectives"}


def parse(text, appendix=False):
    """Returns list of blocks: (kind, payload...)."""
    lines = text.split("\n")
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if re.match(r"^CHAPTER \d$", ln) or re.match(r"^APPENDIX [A-D]$", ln):
            title = lines[i + 1].strip()
            blocks.append(("chapter" if ln.startswith("CHAPTER") else "appendix", ln.strip(), title))
            i += 2
            continue
        if ln.startswith("```"):
            code = []
            i += 1
            while not lines[i].startswith("```"):
                code.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            blocks.append(("code", code))
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        m = re.match(r"^EQ\[(\w+)\]:\s*(.+)$", ln)
        if m:
            blocks.append(("eq", m.group(1), m.group(2)))
            i += 1
            continue
        m = IMG.match(ln)
        if m:
            blocks.append(("fig", m.group(1), m.group(2)))
            i += 1
            continue
        if ln in SUBHEADS:
            blocks.append(("subhead", ln))
            i += 1
            continue
        m = (APPHEAD if appendix else HEAD).match(ln)
        if m and not ln.endswith("."):
            level = ln.split()[0].count(".") + 1
            blocks.append(("heading", level, ln))
            i += 1
            continue
        m = TABCAP.match(ln)
        if m:
            blocks.append(("tabcap", ln))
            i += 1
            continue
        if re.match(r"^\d+\.\s", ln) or ln.startswith("- "):
            marker = re.match(r"^(\d+\.|-)\s", ln).group(1)
            item = ln[len(marker):].strip()
            blocks.append(("item", "•" if marker == "-" else marker, item))
            i += 1
            continue
        blocks.append(("p", ln))
        i += 1
    return blocks


chapter_blocks = [parse(t) for t in chapters]
appendix_blocks = parse(appendix_src, appendix=True)

# entries for pre-filled contents/lists
toc_entries, table_caps, figure_caps, appendix_titles = [], [], [], []
front_titles = ["Declaration", "Recommendation", "Certificate", "Acknowledgements", "Abbreviations/Acronyms",
                "Units and Conversions", "ABSTRACT", "Table of Contents", "List of Tables", "List of Figures",
                "List of Appendices"]
for t in front_titles:
    toc_entries.append((1, t))
for blocks in chapter_blocks:
    for b in blocks:
        if b[0] == "chapter":
            toc_entries.append((1, f"{b[1]}  {b[2]}"))
        elif b[0] == "heading" and b[1] <= 3:
            toc_entries.append((b[1] if b[1] > 1 else 2, b[2]))
        elif b[0] == "tabcap":
            table_caps.append(b[1])
        elif b[0] == "fig":
            figure_caps.append(b[1])
toc_entries.append((1, "REFERENCES"))
for b in appendix_blocks:
    if b[0] == "appendix":
        toc_entries.append((1, f"{b[1].title()}: {b[2].title()}"))
        appendix_titles.append(f"{b[1].title()}: {b[2].title()}")
    elif b[0] == "tabcap":
        table_caps.append(b[1])
    elif b[0] == "fig":
        figure_caps.append(b[1])


def render(blocks, in_appendix=False):
    for b in blocks:
        kind = b[0]
        if kind == "chapter":
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(b[1])
            p.add_run().add_break()
            p.add_run(b[2])
        elif kind == "appendix":
            p = doc.add_paragraph(style="Appendix Heading")
            p.add_run(b[1])
            p.add_run().add_break()
            p.add_run(b[2])
        elif kind == "heading":
            lvl = min(b[1], 4)
            style = {1: "Heading 2", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}[lvl]
            if in_appendix:
                style = "Heading 2"
            para(b[2], style)
        elif kind == "subhead":
            para(b[1], "Subhead")
        elif kind == "p":
            para(b[1])
        elif kind == "item":
            p = doc.add_paragraph(style="List Item")
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
            p.add_run(b[1] + "\t")
            add_text(p, b[2])
        elif kind == "tabcap":
            para(b[1], "Table Caption")
        elif kind == "table":
            add_table(b[1])
        elif kind == "fig":
            add_figure(b[2], b[1])
        elif kind == "eq":
            add_equation(b[2], eq_no[b[1]])
        elif kind == "code":
            for line in b[1]:
                p = doc.add_paragraph(style="Code")
                text = re.sub(r"^ +", lambda m: " " * len(m.group(0)), line)
                p.add_run(text if text else " ")
            doc.add_paragraph(style="Code")


# --------------------------------------------------------------------------------------------
# 5. Front matter
# --------------------------------------------------------------------------------------------
# --- Cover (Appendix 1)
blank(1)
para("Thesis for the Degree of " + DEGREE, "Cover 16")
blank(1)
para(TITLE, "Cover Title")
blank(3)
para("[Logo of Tribhuvan University, 1.6 × 1.6 inch]", "Cover 14")
blank(3)
para(AUTHOR, "Cover 16 Bold")
para(f"({REGNO})", "Cover 16")
para(DEPT, "Cover 14")
blank(1)
para(f"{UNIV}, Kirtipur, Kathmandu, Nepal", "Cover 16 Bold")
blank(1)
para(DATE, "Cover 16 Bold")

# --- Section 2: pretext (roman numerals; title page counted as page i but number not shown)
s_pre = doc.add_section(WD_SECTION.NEW_PAGE)
page_setup(s_pre)
set_pgnum(s_pre, "lowerRoman", 1)
s_pre.different_first_page_header_footer = True
add_footer_page_number(s_pre)
s_pre.first_page_footer.is_linked_to_previous = False
clear_footer_first = s_pre.first_page_footer.paragraphs[0]
# cover section footer: empty
clear_footer(doc.sections[0])

# --- Title page (Appendix 3)
para("Thesis for the Degree of " + DEGREE, "Cover 16")
blank(1)
para(TITLE, "Cover Title")
blank(3)
para(f"Supervised by {SUPERVISOR}", "Cover 16 Bold")
para("A thesis submitted in partial fulfilment of the requirements for the degree of " + DEGREE,
     "Cover 16")
blank(3)
para(AUTHOR, "Cover 16 Bold")
para(f"({REGNO})", "Cover 16")
blank(2)
para(DEPT, "Cover 14")
para(f"{UNIV}, Kirtipur, Kathmandu, Nepal", "Cover 16 Bold")
para(DATE, "Cover 16 Bold")


def sign_block(rows):
    for label in rows:
        para(label, "Front Plain")


# --- Declaration (Appendix 5)
para("Declaration", "Front Heading")
p = para("", "Front Body")
p.add_run("I hereby declare that this study entitled ")
p.add_run(TITLE).bold = True
p.add_run(" is based on my original research work. Related works on the topic by other researchers have been "
          "duly acknowledged. I owe all the liabilities relating to the accuracy and authenticity of the data and "
          "any other information included hereunder.")
blank(2)
sign_block(["Signature: ……………………………", f"Name of the Student: {AUTHOR}",
            f"Registration Number: {REGNO}", "Date: [Date]"])

# --- Recommendation (Appendix 6)
para("Recommendation", "Front Heading")
p = para("", "Front Body")
p.add_run("This is to certify that this thesis entitled ")
p.add_run(TITLE).bold = True
p.add_run(", prepared and submitted by ")
p.add_run(AUTHOR).bold = True
p.add_run(f", in partial fulfilment of the requirements of the degree of {DEGREE} awarded by {UNIV}, has been "
          f"completed under my supervision. I recommend the same for acceptance by {UNIV}.")
blank(2)
sign_block(["Signature: ……………………………", f"Name of the Supervisor: {SUPERVISOR}",
            f"Organization: {DEPT}, {UNIV}", "Date: [Date]"])

# --- Certificate (Appendix 7)
para("Certificate", "Front Heading")
p = para("", "Front Body")
p.add_run("This thesis entitled ")
p.add_run(TITLE).bold = True
p.add_run(" prepared and submitted by ")
p.add_run(AUTHOR).bold = True
p.add_run(f" has been examined by us and is accepted for the award of the degree of {DEGREE} by {UNIV}.")
para("[Print this page on the department letterhead with the official stamp.]", "Front Plain")
for name, role in [("[Name of the external examiner]", "External Examiner"),
                   (f"{SUPERVISOR}", "Supervisor"),
                   ("[Name of the co-supervisor, if applicable]", "Co-supervisor (if applicable)"),
                   ("[Name of the Head of Department]", "Head of Department")]:
    blank(1)
    para(name, "Front Plain")
    para("Signature: ……………………………   Date signed: ……………………", "Front Plain")
    para(role, "Front Plain")

# --- Acknowledgements (Appendix 8)
para("Acknowledgements", "Front Heading")
ack = ("I would like to express my sincere gratitude to my supervisor, Asst. Prof. Jagadish Bhatta, of the Central "
       "Department of Computer Science and Information Technology, Tribhuvan University, for his guidance, "
       "suggestions and encouragement throughout this research.\n"
       "I am thankful to the Head of the Central Department of Computer Science and Information Technology, [name of "
       "Head of Department], and to all the teachers and staff of the department for their support during my "
       "studies. I also thank the creators of the public datasets used in this study for making their data "
       "available, and the providers of the open-source software used in the experiments.\n"
       "I am grateful to my classmates and friends for their discussions and help, and to my family for their "
       "patience, support and encouragement.")
for chunk in ack.split("\n"):
    para(chunk, "Front Body")
blank(1)
sign_block(["Signature: ……………………………", f"Name of the Student: {AUTHOR}", f"Registration Number: {REGNO}",
            "Date: [Date]"])

# --- Abbreviations (Appendix 15)
abbr = [("ALT", "Alanine Aminotransferase"), ("CSIT", "Computer Science and Information Technology"),
        ("CUDA", "Compute Unified Device Architecture"), ("CV", "Cross-Validation"),
        ("EHR", "Electronic Health Record"), ("GPU", "Graphics Processing Unit"),
        ("JIT", "Just-In-Time (compilation)"), ("KNN", "k-Nearest Neighbour"),
        ("MAE", "Mean Absolute Error"), ("M.Sc.", "Master of Science"),
        ("NeurIPS", "Conference on Neural Information Processing Systems"),
        ("ROC-AUC", "Receiver Operating Characteristic – Area Under the Curve"),
        ("SHAP", "Shapley Additive Explanations"), ("SVM", "Support Vector Machine"),
        ("TabFM", "Tabular Foundation Model"), ("TabPFN", "Tabular Prior-Fitted Network"),
        ("TU", "Tribhuvan University")]
abbr.sort(key=lambda x: x[0].lower())
para("List of Abbreviations/Acronyms", "Front Heading")
for a, full in abbr:
    p = doc.add_paragraph(style="Abbrev")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(2.5))
    p.add_run(f"{a}\t{full}")

# --- Units
units = [("%", "Percent"), ("g/dL", "Grams per decilitre"), ("GB", "Gigabyte"), ("mg/dL", "Milligrams per decilitre"),
         ("min", "Minute"), ("ng/mL", "Nanograms per millilitre"), ("s", "Second"), ("U/L", "Units per litre"),
         ("µg/dL", "Micrograms per decilitre"), ("µg/L", "Micrograms per litre")]
para("Units and Conversions", "Front Heading")
for u, meaning in units:
    p = doc.add_paragraph(style="Abbrev")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(2.5))
    p.add_run(f"{u}\t{meaning}")

# --- Abstract (Appendix 9)
pre = open(os.path.join(os.path.dirname(BASE), "drafts", "Pretext_Acknowledgements_to_List_of_Appendices.md"),
           encoding="utf-8").read()
abstract_text = pre.split("**ABSTRACT**", 1)[1].split("Keywords:", 1)[0].strip()
keywords = pre.split("Keywords:", 1)[1].split("\n", 1)[0].strip()
p = doc.add_paragraph(style="Front Heading")
p.add_run("ABSTRACT")
for chunk in abstract_text.split("\n\n"):
    para(chunk.strip(), "Front Body")
blank(1)
p = para("", "Front Plain")
p.add_run("Keywords: ").bold = True
p.add_run(keywords)

# --- Contents and lists (Appendices 10-14): fields, pre-filled, refreshed by Word (Ctrl+A, F9)
para("Table of Contents", "Front Heading")
toc_field(toc_entries, ' TOC \\o "1-3" \\h \\z \\t "Front Heading,1,Appendix Heading,1" ')
para("List of Tables", "Front Heading")
toc_field([(0, t) for t in table_caps], ' TOC \\h \\z \\t "Table Caption,1" ')
para("List of Figures", "Front Heading")
toc_field([(0, t) for t in figure_caps], ' TOC \\h \\z \\t "Figure Caption,1" ')
para("List of Appendices", "Front Heading")
toc_field([(0, t) for t in appendix_titles], ' TOC \\h \\z \\t "Appendix Heading,1" ')

# --------------------------------------------------------------------------------------------
# 6. Main body (arabic page numbers from 1)
# --------------------------------------------------------------------------------------------
s_main = doc.add_section(WD_SECTION.NEW_PAGE)
page_setup(s_main)
set_pgnum(s_main, "decimal", 1)
s_main.different_first_page_header_footer = False
add_footer_page_number(s_main)

for blocks in chapter_blocks:
    render(blocks)

# References (IEEE, numbered by first appearance)
para("REFERENCES", "Front Heading")
for k in order:
    p = doc.add_paragraph(style="Reference")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
    add_text(p, f"[{cite_no[k]}]\t{refs[k]}")
    # Never leave [Accessed: DD Mon. YYYY] unnoticed
render(appendix_blocks, in_appendix=True)

# --------------------------------------------------------------------------------------------
# 7. Settings: update fields when the file is opened, document properties
# --------------------------------------------------------------------------------------------
settings = doc.settings.element
uf = OxmlElement("w:updateFields")
uf.set(qn("w:val"), "true")
anchor = None
for tag in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr",
            "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats", "w:shapeDefaults",
            "w:decimalSymbol", "w:listSeparator"):
    el = settings.find(qn(tag))
    if el is not None:
        anchor = el
        break
if anchor is not None:
    anchor.addprevious(uf)
else:
    settings.append(uf)
doc.core_properties.title = TITLE
doc.core_properties.author = AUTHOR
doc.core_properties.subject = "M.Sc. CSIT thesis, Tribhuvan University"

doc.save(OUT)

# --------------------------------------------------------------------------------------------
# 8. Report
# --------------------------------------------------------------------------------------------
body_words = 0
for blocks in chapter_blocks:
    for b in blocks:
        if b[0] in ("p", "item"):
            body_words += len(resolve(b[-1]).split())
        elif b[0] == "table":
            body_words += sum(len(c.split()) for r in b[1] for c in r)
print("saved", OUT)
print("citations:", len(order), "equations:", len(eq_no), "tables:", len(table_caps), "figures:", len(figure_caps))
print("approx main-body words (chapters 1-5, incl. tables):", body_words)
