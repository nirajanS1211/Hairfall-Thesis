"""Builds Presentation/Hairfall_Presentation.pptx and .pdf (plain black and white, short, with speaker notes)."""
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "figures"
OUT = ROOT / "Presentation"
FONT = "Arial"
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x59, 0x59, 0x59)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 14


def text(slide, x, y, w, h, paras, size=20, bold=False, align=PP_ALIGN.LEFT, color=BLACK, anchor=MSO_ANCHOR.TOP, bullet=False, space=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(paras, str): paras = [paras]
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        r = p.add_run(); r.text = ("•  " + t) if bullet else t
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = color
    return tb


def line(slide, x, y, w):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = BLACK; ln.line.width = Pt(1.25)


def new_slide(title, n, notes):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.6, 0.35, 12.1, 0.9, title, size=32, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    line(s, 0.65, 1.25, 12.0)
    text(s, 0.6, 7.0, 9, 0.35, "Comparing CatBoost, TabPFN, and TabFM for Explainable Hair Fall Risk Prediction", size=11, color=GREY)
    text(s, 11.7, 7.0, 1.0, 0.35, f"{n} / {TOTAL}", size=11, color=GREY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w="12700", cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val="000000")
        etree.SubElement(ln, qn("a:prstDash"), val="solid")


def table(slide, x, y, w, rows, col_w, size=14, header=True, row_h=0.45):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"): tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None: sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # "No Style, Table Grid"
    for j, cw in enumerate(col_w): tbl.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(r):
            c = tbl.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE) if (header and i == 0) else RGBColor(255, 255, 255)
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.04)
            tf = c.text_frame; tf.word_wrap = True; tf.text = ""
            p = tf.paragraphs[0]; run = p.add_run(); run.text = val
            run.font.size = Pt(size); run.font.name = FONT; run.font.color.rgb = BLACK; run.font.bold = (header and i == 0) or (j == 0 and not header)
            cell_border(c)
    return tbl


def picture(slide, name, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), **kw)


# 1 Title -------------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
text(s, 0.8, 0.6, 11.7, 0.5, "Tribhuvan University, Institute of Science and Technology", size=18, align=PP_ALIGN.CENTER, color=GREY)
s.shapes.add_picture(str(FIG / "tu_logo_bw.png"), Inches(6.17), Inches(1.3), height=Inches(1.6))
text(s, 0.8, 3.2, 11.7, 1.7, "Comparing CatBoost, TabPFN, and TabFM for Explainable Hair Fall Risk Prediction", size=36, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
line(s, 3.2, 5.05, 7.0)
text(s, 0.8, 5.25, 11.7, 1.6, ["Dissertation Proposal", "Submitted by: Nirajan Shahi (Roll no. 49/079)", "Supervised by: Asst. Prof. Jagadish Bhatta", "Central Department of Computer Science and Information Technology, Kirtipur, Kathmandu"], size=16, align=PP_ALIGN.CENTER, space=4)
s.notes_slide.notes_text_frame.text = "Introduce the topic: we compare three models for predicting hair fall risk and explain each prediction."

# 2 At a glance ------------------------------------------------------------------------------------
s = new_slide("Project at a glance", 2, "One-slide summary. Problem, aim, data, method, expected result and time. Each point is explained on the next slides.")
table(s, 0.65, 1.5, 12.0, [
    ["Problem", "Hair fall risk is hard to judge, and most prediction models cannot explain their answers."],
    ["Objectives", "1. Benchmark CatBoost, TabPFN, and TabFM.   2. Explain every model with SHAP."],
    ["Research gap", "Foundation models were never tested on hair fall data; CatBoost was never tuned; no side-by-side explanations."],
    ["Data", "21,606 records, 20 features (blood and body measurements, lifestyle, clinical), 3 risk tiers: Low, Moderate, High."],
    ["Method", "Same split for all models; CatBoost tuned, TabPFN and TabFM without tuning; SHAP; accuracy, macro-F1, ROC-AUC, McNemar's test."],
    ["Expected outcome", "A fair comparison, a strong tuned baseline, TabFM limits recorded, and explanations checked against known biology."],
    ["Time", "4 months: proposal work (Months 1 and 2) and final work (Months 3 and 4)."],
], [2.4, 9.6], size=16, header=False, row_h=0.72)

# 3 Problem ---------------------------------------------------------------------------------------
s = new_slide("Problem statement", 3, "Hair loss has many causes that act together, so one test cannot predict it. Doctors rely on history and self-reports. ML can help but is a black box, which is a problem in health care.")
text(s, 0.8, 1.7, 11.7, 4.8, [
    "Hair loss risk is hard to judge: hormones, nutrition, stress, illness, and habits act together.",
    "Assessment often relies on self-reporting or late clinical checks, which are subjective and hard to reach in rural areas.",
    "Earlier prediction studies used classical algorithms and report very different accuracies (about 50% to 100%).",
    "Newer tabular foundation models (TabPFN, TabFM) have not been tested on hair fall data.",
    "Both gradient boosting and foundation models are hard to explain.",
], size=22, bullet=True, space=16)

# 4 Objectives ------------------------------------------------------------------------------------
s = new_slide("Objectives", 4, "Two objectives: compare the three models, and explain them. Everything in the project serves one of these two.")
text(s, 0.8, 1.9, 11.7, 4.5, [
    "1.  To implement and benchmark CatBoost, TabPFN, and TabFM for predicting multi-tier hair fall risk from structured clinical and lifestyle data.",
    "2.  To provide transparent, feature-level decision-support insights for every model through SHAP-based explainability, so that each predicted risk tier can be traced to specific indicators.",
], size=24, space=26)

# 5 Background ------------------------------------------------------------------------------------
s = new_slide("Background: what causes hair loss", 5, "Hair grows in cycles. Anything that pushes follicles out of the growth phase early causes shedding. These factors overlap, which is why a model that combines many indicators is useful.")
table(s, 0.65, 1.5, 12.0, [
    ["Factor", "What it does", "Source"],
    ["Hormones and heredity", "Inherited sensitivity to androgens; family history is a key indicator", "Agaoglu et al. (2021)"],
    ["Nutrition and iron", "Shortage of iron, zinc, and vitamins weakens follicles", "Guo & Katta (2017); Lin et al. (2023)"],
    ["Thyroid function", "Thyroid hormones regulate the hair cycle", "Hussein et al. (2023)"],
    ["Psychological stress", "Pushes follicles into the resting phase early", "Bai et al. (2026)"],
    ["Illness and infection", "Can trigger heavy shedding (telogen effluvium)", "Cline et al. (2021)"],
], [3.2, 5.6, 3.2], size=16, row_h=0.6)
text(s, 0.65, 5.3, 12, 1.2, "Trend: hair loss is reported more often and at a younger age, and it affects self-confidence and quality of life.", size=18)

# 6 Literature -------------------------------------------------------------------------------------
s = new_slide("Literature: what others did", 6, "These are the closest studies. Most use classical models on survey data, or images. Results differ a lot between studies, and explanations were not compared.")
table(s, 0.5, 1.45, 12.3, [
    ["Study", "Data", "Models", "Main result"],
    ["Khatun et al. (2022)", "Survey, 610 people", "SVM, KNN, LR, RF, XGBoost", "XGBoost 92.62%"],
    ["Sai et al. (2023)", "Hair fall data", "SVM, KNN, DT, RF, LR, ensemble", "Ensemble best"],
    ["Kumar et al. (2025)", "2,000 records", "RF, XGBoost, CatBoost, LightGBM", "RF 100%, CatBoost 49.5%"],
    ["Siami & Azis (2025)", "Multi-factor data", "LR, DT, RF, GB, XGBoost, voting", "About 50% at best"],
    ["Leema et al. (2025)", "Survey, 750 people", "LSTM, RF, TFT, ARIMAX", "TFT 97.5%"],
    ["Shakeel et al. (2021); Sayyad et al. (2022)", "Hair images", "SVM, KNN, VGG with SVM", "91.4% and 98.31%"],
], [3.9, 2.6, 3.6, 2.2], size=14, row_h=0.62)
text(s, 0.5, 6.0, 12.3, 0.8, "LR = Logistic Regression, DT = Decision Tree, RF = Random Forest, GB = Gradient Boosting", size=12, color=GREY)

# 7 Models literature -----------------------------------------------------------------------------
s = new_slide("Literature: the three model families", 7, "CatBoost is a strong boosted-tree model for categorical data. TabPFN and TabFM are new foundation models: pretrained once, then they predict in one pass without training.")
text(s, 0.8, 1.6, 11.7, 5.0, [
    "CatBoost (Prokhorenkova et al., 2018): gradient boosted trees with ordered encoding of categorical answers.",
    "TabPFN (Hollmann et al., 2025): transformer pretrained on synthetic tables; predicts in one forward pass; best on small tables.",
    "TabFM (Kong et al., 2026): 400-million-parameter foundation model; first among default foundation models on 51 TabArena datasets.",
    "An independent check of TabFM reported software defects and memory limits on larger tables (Pandey, 2026).",
    "SHAP (Lundberg & Lee, 2017): explains any model by the contribution of each feature.",
], size=21, bullet=True, space=16)

# 8 Research gap ----------------------------------------------------------------------------------
s = new_slide("Research gap", 8, "This is what we found missing in the literature. It is why the project compares these three models and explains them.")
text(s, 0.8, 1.7, 11.7, 4.8, [
    "Hair loss prediction studies used only classical algorithms, and accuracies differ widely.",
    "CatBoost was compared only once, and it was not tuned.",
    "TabPFN and TabFM have not been tested on hair loss data.",
    "TabFM has only one independent check, which reported defects and memory limits.",
    "No study explains a boosted-tree model and foundation models side by side.",
], size=24, bullet=True, space=20)

# 9 Data -------------------------------------------------------------------------------------------
s = new_slide("Data", 9, "Two public datasets are the sources. A super dataset of 200,000 records holds all their fields plus extra fields. Age and gender are identified in it. Comparing it with the two datasets gives the final dataset of 21,606 records used for modelling.")
text(s, 0.6, 1.5, 5.0, 5.3, [
    "Dataset 1 (Kaggle): 100,000 records, 13 columns of measurements and hair fall (0 to 5).",
    "Dataset 2 (Mendeley): survey of 716 people, 14 columns.",
    "Super dataset: 200,000 records with all fields of both and extra fields.",
    "Final dataset: 21,606 records, 20 features, target Low 45%, Moderate 35%, High 20%.",
], size=17, bullet=True, space=12)
picture(s, "fig_3_2_dataset_link.png", 5.7, 2.4, w=7.3)

# 10 Framework -------------------------------------------------------------------------------------
s = new_slide("Research framework", 10, "One pipeline: clean and split the data once, give the same rows to the three models, evaluate them the same way, and explain each with SHAP. The 500 and 2,000 record sets show behaviour with little data.")
picture(s, "fig_3_1_framework.png", 1.9, 1.5, w=9.5)
text(s, 0.9, 6.3, 11.5, 0.6, "Same split and seed for all models. Training sizes: 500, 2,000, and all 17,284 records.", size=16, align=PP_ALIGN.CENTER)

# 11 Models ----------------------------------------------------------------------------------------
s = new_slide("Models", 11, "Three different ways to predict. CatBoost needs training and tuning. TabPFN and TabFM only read the training rows as context and predict, so they need no tuning.")
table(s, 0.65, 1.6, 12.0, [
    ["Model", "How it works", "Training and tuning"],
    ["CatBoost", "Many small symmetric decision trees, each correcting the previous ones", "Trained and tuned: grid search with 5-fold cross-validation"],
    ["TabPFN", "Transformer reads training rows and the patient, predicts in one pass", "None; weights are fixed"],
    ["TabFM", "Transformer with column and row attention, about 400 million parameters", "None; weights are fixed"],
], [2.2, 6.0, 3.8], size=16, row_h=1.0)
text(s, 0.65, 5.9, 12, 0.8, "TabPFN and TabFM will run on a Kaggle GPU; CatBoost and the statistics on a Mac.", size=16)

# 12 Explain + evaluate ---------------------------------------------------------------------------
s = new_slide("Explainability and evaluation", 12, "SHAP shows which features push each prediction up or down. Metrics are macro averages because the tiers are not equal in size. McNemar's test tells us if differences are real.")
text(s, 0.6, 1.6, 5.9, 5.0, [
    "SHAP for all three models:",
    "TreeExplainer for CatBoost",
    "KernelExplainer for TabPFN and TabFM",
    "Global ranking and per-patient explanations",
], size=20, bullet=False, space=12)
text(s, 6.9, 1.6, 5.9, 5.0, [
    "Evaluation on the same test set:",
    "Accuracy, macro-precision, macro-recall, macro-F1",
    "ROC-AUC (one-versus-rest)",
    "Confusion matrix",
    "McNemar's test and Wilson 95% intervals",
], size=20, space=12)

# 13 Expected outcomes -----------------------------------------------------------------------------
s = new_slide("Expected outcomes", 13, "These are expectations to test, not results. The study can also show the opposite, and a difference counts only if McNemar's test supports it.")
text(s, 0.8, 1.7, 11.7, 4.8, [
    "A fair comparison of three model families on the same hair fall data.",
    "A properly tuned CatBoost as a strong baseline.",
    "TabPFN and TabFM competitive without tuning, especially with small training sets (to be tested).",
    "The practical limits of TabFM recorded on a real health dataset.",
    "SHAP explanations that agree with known causes: iron, stress, family history.",
], size=22, bullet=True, space=16)

# 14 Schedule -------------------------------------------------------------------------------------
s = new_slide("Schedule: 4 months", 14, "Months 1 and 2 are the proposal work. Months 3 and 4 are the final work after approval, in the order of the methodology. Thank you; questions are welcome.")
picture(s, "fig_4_1_gantt.png", 2.6, 1.4, h=5.4)

prs.save(OUT / "Hairfall_Presentation.pptx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(OUT / "Hairfall_Presentation.pptx")], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print("ok", len(prs.slides.__iter__.__self__._sldIdLst))
