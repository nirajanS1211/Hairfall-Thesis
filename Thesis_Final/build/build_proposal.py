"""Builds Hairfall_Proposal.docx and .pdf from the markdown chapters in Thesis_Final.
Format follows FORMAT_RULES.md: Times New Roman 12 pt, 1.5 spacing, 1-inch margins, one space after each paragraph,
chapter heading 16 pt bold, section 14 pt bold, sub-section 12 pt bold, equations numbered (i), (ii), ...
Run:  python build/build_proposal.py
"""
import re, subprocess, sys
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor, Emu

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / "build"
OUT_DOCX = ROOT / "Hairfall_Proposal.docx"
CHAPTERS = ["Chapter_1_Introduction.md", "Chapter_2_Background_and_Literature_Review.md",
            "Chapter_3_Methodology.md", "Chapter_4_Expected_Outcome_and_Schedule.md"]
FONT = "Times New Roman"
TEXT_W = 6.27  # A4 width 8.27 in minus two 1-inch margins
ROMAN = ["i","ii","iii","iv","v","vi","vii","viii","ix","x","xi","xii","xiii","xiv","xv"]

plt.rcParams["mathtext.fontset"] = "stix"
plt.rcParams["font.family"] = "STIXGeneral"

# ---------------------------------------------------------------- equations
def latex_for_mathtext(s):
    s = re.sub(r"\\tag\{[^}]*\}", "", s)
    s = s.replace("\\big(", "(").replace("\\big)", ")").replace("\\Big(", "(").replace("\\Big)", ")")
    s = s.replace("\\dfrac", "\\frac")
    def text_repl(m):
        return "\\mathrm{" + m.group(1).replace(" ", "\\ ").replace("-", "\\text{-}") + "}"
    s = re.sub(r"\\text\{([^}]*)\}", lambda m: "\\mathrm{" + m.group(1).replace("-", " ").replace(" ", "\\ ") + "}", s)
    return s.strip()

def render_equation(latex, n):
    path = BUILD / f"eq_{n}.png"
    fig = plt.figure(figsize=(0.1, 0.1))
    fig.text(0, 0, f"${latex_for_mathtext(latex)}$", fontsize=14)
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=True)
    plt.close(fig)
    from PIL import Image
    w, h = Image.open(path).size
    return path, w / 300, h / 300

# ---------------------------------------------------------------- docx helpers
def set_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)

def fmt(p, align=None, space_after=12, spacing=1.5, before=0, keep_next=False, left=None, first=None, hanging=None):
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(before)
    pf.line_spacing = spacing
    if align is not None: p.alignment = align
    if keep_next: pf.keep_with_next = True
    if left is not None: pf.left_indent = Inches(left)
    if first is not None: pf.first_line_indent = Inches(first)
    if hanging is not None:
        pf.left_indent = Inches(hanging); pf.first_line_indent = Inches(-hanging)

SUB_LEFT = {"D", "x", "h", "F", "y", "s", "f", "TP", "FP", "FN", "Precision", "Recall", "F1", "AUC", "p"}
SUB_RE = re.compile(r"(?<![A-Za-z0-9_])([A-Za-zφ][A-Za-z0-9]*)_(\([^)]+\)|[A-Za-zθ0-9]{1,5})(?![A-Za-z0-9_])")

def add_text(p, text, size=12, bold=False, italic=False):
    """Adds text with **bold**, *italic*, subscripts (x_k) and yellow [CONFIRM ...] highlights."""
    text = text.replace("φⱼ", "φ_j")
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*", text):
        if m.start() > pos: _plain(p, text[pos:m.start()], size, bold, italic)
        if m.group(1) is not None: _plain(p, m.group(1), size, True, italic)
        else: _plain(p, m.group(2), size, bold, True)
        pos = m.end()
    if pos < len(text): _plain(p, text[pos:], size, bold, italic)

def _plain(p, s, size, bold, italic):
    pos = 0
    def emit(t, sub=False, hl=False):
        if not t: return
        r = p.add_run(t); set_font(r, size, bold, italic)
        if sub: r.font.subscript = True
        if hl: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for chunk in re.split(r"(\[CONFIRM:[^\]]*\])", s):
        if chunk.startswith("[CONFIRM:"):
            emit(chunk, hl=True); continue
        pos = 0
        for m in SUB_RE.finditer(chunk):
            left, right = m.group(1), m.group(2)
            if left in ("d1", "d2") or not (len(left) <= 2 or left in SUB_LEFT): continue
            if right.startswith("("): right = right[1:-1]
            emit(chunk[pos:m.start()]); emit(left); emit(right, sub=True); pos = m.end()
        emit(chunk[pos:])

def add_page_field(paragraph):
    r = paragraph.add_run(); set_font(r, 12)
    for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if typ:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), typ)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)

def set_pgnum(section, fmt_, start):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")): sectPr.remove(old)
    e = OxmlElement("w:pgNumType"); e.set(qn("w:fmt"), fmt_); e.set(qn("w:start"), str(start)); sectPr.append(e)

def shade(cell, color="E7E6E6"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def style_setup(doc):
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        h = doc.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
        h.font.italic = False; h.font.color.rgb = RGBColor(0, 0, 0)
        rpr = h.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), FONT)
        h.paragraph_format.keep_with_next = True

# ---------------------------------------------------------------- markdown -> docx
class Builder:
    def __init__(self, doc):
        self.doc = doc; self.eq_n = 0

    def heading(self, text, level, page_break=False):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        r = p.add_run(text); set_font(r, {1: 16, 2: 14, 3: 12}[level], True)
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT, space_after=12 if level > 1 else 18, spacing=1.5, before=6 if level > 1 else 0, keep_next=True)
        if page_break: p.paragraph_format.page_break_before = True
        return p

    def para(self, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, **kw):
        p = self.doc.add_paragraph(); add_text(p, text); fmt(p, align, **kw); return p

    def caption(self, text, above):
        p = self.doc.add_paragraph()
        m = re.match(r"\*\*(Figure|Table) (\d\.\d):\*\*\s*(.*)", text)
        add_text(p, f"**{m.group(1)} {m.group(2)}:** {m.group(3)}")
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, space_after=12 if not above else 6, spacing=1.5, keep_next=above)
        return p

    def image(self, path):
        from PIL import Image
        w, h = Image.open(path).size
        width = TEXT_W if w > 1400 else min(TEXT_W, w / 220)
        p = self.doc.add_paragraph(); p.add_run().add_picture(str(path), width=Inches(width))
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, space_after=6, spacing=1.0, keep_next=True)

    def equation(self, latex):
        """Equation slightly left of centre; a dotted leader starts where the equation ends and runs to the number (i), (ii), ..."""
        from docx.enum.table import WD_ALIGN_VERTICAL
        tag = re.search(r"\\tag\{([^}]*)\}", latex).group(1)
        self.eq_n += 1
        path, w, h = render_equation(latex, self.eq_n)
        total, numw = 5.5, 1.1
        w = min(w, total - numw - 0.5)
        spacer = max(0.1, (total - numw - w) / 2)
        widths = [spacer, w + 0.3, total - spacer - w - 0.3]
        t = self.doc.add_table(rows=1, cols=3); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for col, wd in zip(t.columns, widths): col.width = Inches(wd)
        for c, wd in zip(t.rows[0].cells, widths):
            c.width = Inches(wd); c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        pm = t.rows[0].cells[1].paragraphs[0]
        pm.add_run().add_picture(str(path), width=Inches(w))
        fmt(pm, WD_ALIGN_PARAGRAPH.LEFT, space_after=0, spacing=1.0)
        pn = t.rows[0].cells[2].paragraphs[0]
        pn.paragraph_format.tab_stops.add_tab_stop(Inches(widths[2] - 0.15), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = pn.add_run("\t(" + tag + ")"); set_font(r, 12)
        fmt(pn, WD_ALIGN_PARAGRAPH.LEFT, space_after=0, spacing=1.0)
        fmt(t.rows[0].cells[0].paragraphs[0], None, space_after=0, spacing=1.0)
        sp = self.doc.add_paragraph(); fmt(sp, None, space_after=6, spacing=1.0)

    def table(self, rows):
        ncol = len(rows[0]); size = 12 if ncol <= 4 else 9.5
        t = self.doc.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        weights = [min(max(max(len(r[c]) for r in rows[1:]), len(rows[0][c]) // 2), 30) + 4 for c in range(ncol)]
        widths = [TEXT_W * w / sum(weights) for w in weights]
        for col, wd in zip(t.columns, widths): col.width = Inches(wd)
        for i, row in enumerate(rows):
            tr = t.rows[i]
            trPr = tr._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            if i == 0:
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            for j, val in enumerate(row):
                c = tr.cells[j]; c.width = Inches(widths[j])
                p = c.paragraphs[0]; add_text(p, val, size=size, bold=(i == 0))
                fmt(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=2, spacing=1.0, before=2)
                if i == 0: shade(c)
        sp = self.doc.add_paragraph(); fmt(sp, None, space_after=6, spacing=1.0)

    def list_item(self, text, bullet, num=None):
        p = self.doc.add_paragraph()
        lead = "•\t" if bullet else f"{num}.\t"
        r = p.add_run(lead); set_font(r, 12)
        add_text(p, text)
        fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8, spacing=1.5)
        p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))

    def markdown(self, path, skip_first_h1=False):
        lines = Path(path).read_text().splitlines()
        i = 0; first = True; num = 0
        while i < len(lines):
            ln = lines[i].rstrip(); i += 1
            if not ln.strip(): num = num if re.match(r"\d+\. ", lines[i] if i < len(lines) else "") else 0; continue
            if ln.startswith("# "):
                self.heading(ln[2:], 1, page_break=True); continue
            if ln.startswith("## "): self.heading(ln[3:], 2); continue
            if ln.startswith("### "): self.heading(ln[4:], 3); continue
            m = re.match(r"!\[.*?\]\((.*?)\)", ln)
            if m: self.image(Path(path).parent / m.group(1)); continue
            if ln.startswith("$$"):
                buf = ln
                while not buf.rstrip().endswith("$$") or buf.strip() == "$$":
                    buf += lines[i]; i += 1
                self.equation(buf.strip().strip("$")); continue
            if ln.startswith("|"):
                rows = []
                i -= 1
                while i < len(lines) and lines[i].startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells): rows.append(cells)
                    i += 1
                self.table(rows); continue
            if re.match(r"\*\*(Figure|Table) \d\.\d:\*\*", ln):
                self.caption(ln, above=ln.startswith("**Table")); continue
            if ln.startswith("- "): self.list_item(ln[2:], True); continue
            m = re.match(r"(\d+)\. (.*)", ln)
            if m and not ln.startswith("**"): self.list_item(m.group(2), False, m.group(1)); continue
            self.para(ln)

# ---------------------------------------------------------------- document parts
def title_page(doc):
    """Same first page as 'Hair fall Proposal Revised (Updated)': 14 pt bold, centred, TU logo."""
    def line(text="", after=6):
        p = doc.add_paragraph(); r = p.add_run(text if text else " "); set_font(r, 14, True)
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, space_after=after, spacing=1.08)
    line("Tribhuvan University"); line("Institute of Science and Technology"); line(); line()
    p = doc.add_paragraph(); p.add_run().add_picture(str(ROOT / "figures" / "tu_logo.jpg"), width=Inches(1.39))
    fmt(p, WD_ALIGN_PARAGRAPH.CENTER, space_after=3, spacing=1.08)
    line(); line(); line("A Dissertation Proposal on", 6)
    line("\u201cComparing CatBoost, TabPFN, and TabFM for Explainable Hair Fall Risk Prediction\u201d", 6)
    line(); line("Supervised by", 6); line("Asst. Prof. Jagadish Bhatta", 6); line()
    line("Submitted by", 6); line("Nirajan Shahi", 6); line("Roll no. 49/079", 6); line(); line()
    line("Submitted to", 6); line("Central Department of Computer Science and Information Technology", 6)
    line("Tribhuvan University, Kirtipur", 6); line("Kathmandu, Nepal", 6); line()
    line("In partial fulfillment of the requirement for Master\u2019s Degree in Computer Science and Information technology (M.Sc. CSIT)", 6)

def front_title(doc, text, first=False):
    p = doc.add_paragraph(); r = p.add_run(text); set_font(r, 16, True)
    fmt(p, WD_ALIGN_PARAGRAPH.CENTER, space_after=14, spacing=1.5, keep_next=True)
    if not first: p.paragraph_format.page_break_before = True

def leader_line(doc, text, page, indent=0.0, bold=False, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(text); set_font(r, 12, bold)
    r = p.add_run("\t" + str(page)); set_font(r, 12, bold)
    fmt(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=after, spacing=1.5, left=indent)
    p.paragraph_format.right_indent = Inches(0.0)

def parse_front():
    txt = (ROOT / "Front_Matter_Lists.md").read_text()
    abbr = [tuple(c.strip() for c in l.strip().strip("|").split("|")) for l in txt.split("# List of Abbreviations")[1].splitlines()
            if l.startswith("|") and not l.startswith("|---") and not l.startswith("| |")]
    return abbr

def collect_entries():
    heads, figs, tabs = [], [], []
    for f in CHAPTERS + ["References.md"]:
        for l in (ROOT / f).read_text().splitlines():
            m = re.match(r"(#{1,3}) (.*)", l)
            if m: heads.append((len(m.group(1)), m.group(2)))
            m = re.match(r"\*\*(Figure|Table) (\d\.\d):\*\*\s*(.*)", l)
            if m: (figs if m.group(1) == "Figure" else tabs).append((f"{m.group(1)} {m.group(2)}", m.group(3)))
    return heads, figs, tabs

def build(pages):
    heads, figs, tabs = collect_entries()
    doc = Document(); style_setup(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, a, Inches(1))
    title_page(doc)
    # front matter section
    s2 = doc.add_section(WD_SECTION.NEW_PAGE); set_pgnum(s2, "lowerRoman", 2)
    s2.footer.is_linked_to_previous = False
    fp = s2.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_page_field(fp)
    front_title(doc, "Table of Contents", first=True)
    for lvl, text in heads:
        key = text
        pg = pages.get(("h", key), "")
        if lvl == 1: leader_line(doc, text, pg, 0.0, bold=True, after=4)
        elif lvl == 2: leader_line(doc, text, pg, 0.3)
        else: leader_line(doc, text, pg, 0.7)
    front_title(doc, "List of Figures")
    for lab, cap in figs: leader_line(doc, f"{lab}: {cap}", pages.get(("c", lab), ""), 0.0, after=6)
    front_title(doc, "List of Tables")
    for lab, cap in tabs: leader_line(doc, f"{lab}: {cap}", pages.get(("c", lab), ""), 0.0, after=6)
    front_title(doc, "List of Abbreviations")
    for a, b in parse_front():
        p = doc.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Inches(1.4))
        r = p.add_run(a + "\t" + b); set_font(r, 12)
        fmt(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=4, spacing=1.5, hanging=1.4)
    # body section
    s3 = doc.add_section(WD_SECTION.NEW_PAGE); set_pgnum(s3, "decimal", 1)
    s3.footer.is_linked_to_previous = False
    fp = s3.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_page_field(fp)
    b = Builder(doc)
    for i, f in enumerate(CHAPTERS):
        b.markdown(ROOT / f)
    # references
    ref = (ROOT / "References.md").read_text().splitlines()
    b.heading("References", 1, page_break=True)
    for l in ref:
        if not l.strip() or l.startswith("#") or l.startswith("*(APA"): continue
        p = doc.add_paragraph(); add_text(p, l)
        fmt(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=8, spacing=1.5, hanging=0.5)
    doc.save(OUT_DOCX)

def to_pdf():
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(ROOT), str(OUT_DOCX)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return ROOT / (OUT_DOCX.stem + ".pdf")

def find_pages(pdf):
    from pypdf import PdfReader
    texts = [re.sub(r"^(?:[ivx]+|\d+) ", "", re.sub(r"\s+", " ", (p.extract_text() or "")).strip()) for p in PdfReader(str(pdf)).pages]
    heads, figs, tabs = collect_entries()
    # body starts at the first page (after the front matter) that begins with the first chapter heading
    body0 = next(i for i, t in enumerate(texts) if t.strip().startswith(heads[0][1]) and "Table of Contents" not in t[:40] and i > 1)
    pages, cur = {}, body0
    for lvl, text in heads:
        norm = re.sub(r"\s+", " ", text)
        for i in range(cur, len(texts)):
            if norm in texts[i]: pages[("h", text)] = i - body0 + 1; cur = i; break
    for lab, cap in figs + tabs:
        for i in range(body0, len(texts)):
            if f"{lab}: {cap[:25]}" in texts[i]: pages[("c", lab)] = i - body0 + 1; break
    return pages, len(texts), body0

if __name__ == "__main__":
    build({})
    pdf = to_pdf(); pages, n, body0 = find_pages(pdf)
    build(pages); pdf = to_pdf(); pages2, n2, body02 = find_pages(pdf)
    print("pages", n2, "body starts at pdf page", body02 + 1, "| stable:", pages == pages2)
    missing = [k for k in pages2 if pages2[k] == ""]
    print("entries with page numbers:", len(pages2))
