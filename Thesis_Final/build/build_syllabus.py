"""Builds Syllabus_Dissertation_CSc666.docx and .pdf from the markdown file (Times New Roman 12 pt, 1.5 spacing, 1-inch margins)."""
import re, subprocess
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
ROOT = Path(__file__).resolve().parent.parent
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
s = doc.sections[0]; s.page_width, s.page_height = Inches(8.27), Inches(11.69)
for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(s, a, Inches(1))
def runs(p, text, size=12, bold=False):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part: continue
        r = p.add_run(part); r.font.name = "Times New Roman"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(size); r.bold = bold or (i % 2 == 1); r.font.color.rgb = RGBColor(0, 0, 0)
for ln in (ROOT / "Syllabus_Dissertation_CSc666.md").read_text().splitlines():
    if not ln.strip(): continue
    p = doc.add_paragraph(); pf = p.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(12)
    if ln.startswith("# "):
        runs(p, ln[2:], 16, True); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith("## "):
        runs(p, ln[3:], 14, True); pf.keep_with_next = True
    elif ln.startswith("    - "):
        runs(p, ln[6:]); pf.left_indent = Inches(0.9); pf.first_line_indent = Inches(-0.3); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif re.match(r"\d+\. ", ln):
        m = re.match(r"(\d+)\. (.*)", ln); runs(p, f"{m.group(1)}. " + m.group(2)); pf.left_indent = Inches(0.5); pf.first_line_indent = Inches(-0.4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        runs(p, ln); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if ln.startswith("**") and ln.count("**") == 2 and len(ln) < 60: pf.space_after = Pt(2); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
out = ROOT / "Syllabus_Dissertation_CSc666.docx"; doc.save(out)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(ROOT), str(out)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print("ok")
